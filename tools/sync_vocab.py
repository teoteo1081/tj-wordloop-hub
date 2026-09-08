"""
================================================================================
 sync_vocab.py — KÉO/ĐẨY TỪ VỰNG QUA LẠI GIỮA SUPABASE VÀ EXCEL
================================================================================
 BẢN 2026-09-07 (v1)

 2 CHIỀU, chọn bằng MODE bên dưới:

   MODE = "export"  ->  Kéo TOÀN BỘ từ vựng (kèm ngữ cảnh Hub/Notebook/
                         Section/Page/Batch/Block cho dễ nhìn) từ Supabase
                         xuống 1 file Excel để xem/sửa tay thoải mái.

   MODE = "import"  ->  Đọc LẠI file Excel đó (đã sửa xong), đẩy NGƯỢC lên
                         Supabase. Ghép theo cột "id" (word id):
                           - Có "id" khớp từ cũ  -> CẬP NHẬT từ đó (ghi đè
                             term/level/pos/ipa/def_en/meaning_vi/sort).
                           - "id" ĐỂ TRỐNG        -> TẠO TỪ MỚI (bắt buộc
                             phải điền "block_id" - biết thêm vào Block
                             nào - lấy từ đúng cột "block_id" của 1 dòng
                             khác trong CÙNG Block đó, copy xuống).
                         XOÁ HẲN 1 dòng khỏi Excel KHÔNG xoá từ đó trên
                         Supabase (an toàn - tool không tự xoá dữ liệu,
                         muốn xoá từ thật thì tự vào Supabase Table Editor).

 QUY TRÌNH DÙNG:
   1. Sửa MODE = "export" bên dưới, chạy: python tools/sync_vocab.py
      -> ra file Excel (đường dẫn in cuối cùng, mặc định tools/reports/).
   2. Mở Excel, sửa tay (đổi nghĩa, thêm dòng mới cho từ mới...).
   3. Sửa MODE = "import", chạy lại: python tools/sync_vocab.py
      -> đẩy thẳng lên Supabase.
   4. Đổi lại MODE = "export" và chạy lại bất cứ lúc nào để lấy bản MỚI
      NHẤT (VD sau khi có người học thêm/sửa qua chính app).

 ⚠️ File Excel này CHỈ chứa nội dung học (không phải link đăng nhập như
 manage_users.py) nên không nhạy cảm - nhưng vẫn lưu cục bộ (tools/reports/,
 đã có trong .gitignore), không cần đẩy lên GitHub.
================================================================================
"""

import datetime
import json
import re
import sys
import uuid
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ==============================================================================
# CONFIG — SỬA CHỖ NÀY
# ==============================================================================

MODE = "export"   # "export" (Supabase -> Excel) hoặc "import" (Excel -> Supabase)

EXCEL_FILE = Path(__file__).resolve().parent / "reports" / "tu_vung.xlsx"

# import mode: số dòng gửi lên Supabase mỗi lượt (PostgREST upsert theo lô,
# tránh gửi 1 lần cả ~7000 dòng dễ bị timeout/lỗi payload quá lớn).
IMPORT_BATCH_SIZE = 300

AUTO_OPEN_EXCEL = True

# ==============================================================================
# LẤY URL + ANON KEY TỪ js/config.js (y hệt manage_users.py)
# ==============================================================================

CONFIG_JS_PATH = Path(__file__).resolve().parent.parent / "js" / "config.js"


def _read_config_js():
    text = CONFIG_JS_PATH.read_text(encoding="utf-8")
    url_m = re.search(r'SUPABASE_URL:\s*"([^"]+)"', text)
    key_m = re.search(r'SUPABASE_ANON_KEY:\s*"([^"]+)"', text)
    if not url_m or not key_m:
        sys.exit(f"❌ Không đọc được SUPABASE_URL/SUPABASE_ANON_KEY từ {CONFIG_JS_PATH}")
    return url_m.group(1).rstrip("/"), key_m.group(1)


SUPABASE_URL, SUPABASE_ANON_KEY = _read_config_js()
REST_URL = SUPABASE_URL + "/rest/v1"
HEADERS = {
    "apikey": SUPABASE_ANON_KEY,
    "Authorization": f"Bearer {SUPABASE_ANON_KEY}",
    "Content-Type": "application/json",
}

WORD_COLS = ["id", "block_id", "sort", "term", "level", "pos", "ipa", "def_en", "meaning_vi"]


def fetch_all(table, select, page_size=1000):
    rows = []
    offset = 0
    while True:
        headers = dict(HEADERS)
        headers["Range-Unit"] = "items"
        headers["Range"] = f"{offset}-{offset + page_size - 1}"
        r = requests.get(f"{REST_URL}/{table}", headers=headers,
                          params={"select": select}, timeout=30)
        r.raise_for_status()
        chunk = r.json()
        rows.extend(chunk)
        if len(chunk) < page_size:
            break
        offset += page_size
    return rows


# ==============================================================================
# EXPORT: Supabase -> Excel
# ==============================================================================

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
ID_FONT = Font(name="Arial", size=9, color="999999")


def do_export():
    print("Đang tải dữ liệu từ Supabase (hubs -> ... -> words)...")
    hubs = {h["id"]: h for h in fetch_all("hubs", "id,name")}
    notebooks = {n["id"]: n for n in fetch_all("notebooks", "id,hub_id,name")}
    sections = {s["id"]: s for s in fetch_all("sections", "id,notebook_id,name")}
    pages = {p["id"]: p for p in fetch_all("pages", "id,section_id,name")}
    batches = {b["id"]: b for b in fetch_all("batches", "id,page_id,name")}
    blocks = {b["id"]: b for b in fetch_all(
        "blocks", "id,batch_id,name,sort,context_passage,context_passage_candidates")}
    words = fetch_all("words", ",".join(WORD_COLS))
    print(f"  -> {len(words)} từ, {len(blocks)} block, {len(batches)} batch, "
          f"{len(pages)} page, {len(sections)} section, {len(notebooks)} notebook, {len(hubs)} hub.")

    def _chain_of(block_id):
        """Trả về (hub_name, notebook_name, section_name, page_name, batch_name, block_name)
        cho 1 block - đi ngược lên hết chuỗi cha, để trống nếu đứt đoạn ở
        đâu đó (dữ liệu lỗi/mồ côi) thay vì crash cả script."""
        blk = blocks.get(block_id, {})
        bat = batches.get(blk.get("batch_id"), {})
        pg = pages.get(bat.get("page_id"), {})
        sec = sections.get(pg.get("section_id"), {})
        nb = notebooks.get(sec.get("notebook_id"), {})
        hub = hubs.get(nb.get("hub_id"), {})
        return (hub.get("name", ""), nb.get("name", ""), sec.get("name", ""),
                pg.get("name", ""), bat.get("name", ""), blk.get("name", ""))

    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Từ vựng"

    headers = ["id (ĐỂ TRỐNG = tạo từ mới)", "block_id (bắt buộc nếu tạo mới)",
               "Hub", "Notebook", "Section", "Page", "Batch", "Block",
               "sort", "term", "level", "pos", "ipa", "def_en", "meaning_vi"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Sort theo đúng thứ tự cây (hub > notebook > ... > block > sort trong
    # block) cho dễ đọc/đối chiếu, thay vì thứ tự ngẫu nhiên Supabase trả về.
    def _sort_key(w):
        chain = _chain_of(w.get("block_id"))
        return chain + (w.get("sort") or 0,)

    words.sort(key=_sort_key)

    for r_idx, w in enumerate(words, start=2):
        chain = _chain_of(w.get("block_id"))
        c_id = ws.cell(row=r_idx, column=1, value=w.get("id"))
        c_id.font = ID_FONT
        c_bid = ws.cell(row=r_idx, column=2, value=w.get("block_id"))
        c_bid.font = ID_FONT
        for i, v in enumerate(chain, start=3):
            ws.cell(row=r_idx, column=i, value=v)
        ws.cell(row=r_idx, column=9, value=w.get("sort"))
        ws.cell(row=r_idx, column=10, value=w.get("term"))
        ws.cell(row=r_idx, column=11, value=w.get("level"))
        ws.cell(row=r_idx, column=12, value=w.get("pos"))
        ws.cell(row=r_idx, column=13, value=w.get("ipa"))
        ws.cell(row=r_idx, column=14, value=w.get("def_en"))
        ws.cell(row=r_idx, column=15, value=w.get("meaning_vi"))

    widths = [26, 26, 16, 16, 16, 20, 12, 12, 6, 22, 8, 8, 16, 40, 30]
    for i, wd in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(words) + 1}"

    # Sheet "Bài đọc" - context_passage/context_passage_candidates của TỪNG
    # BLOCK (không phải từng từ - 1 Block có 1 dòng ở đây, khác sheet "Từ
    # vựng" là 1 dòng/từ). context_passage_candidates là JSONB (mảng bài
    # đọc ứng cử, tối đa 3 bài/Block theo quy trình viết bài đọc hiện tại,
    # xem CLAUDE.md) - hiển thị dạng chuỗi JSON để xem/sửa tay được trong
    # Excel, import lại sẽ tự parse ngược lại.
    ws3 = wb.create_sheet("Bài đọc")
    ba_headers = ["id (block id - GIỮ NGUYÊN)", "Hub", "Notebook", "Section", "Page",
                  "Batch", "Block", "context_passage (bài đọc chính, nếu có)",
                  "context_passage_candidates (JSON - danh sách bài đọc ứng cử)"]
    for i, h in enumerate(ba_headers, start=1):
        c = ws3.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    block_list = sorted(blocks.values(), key=lambda b: _chain_of(b["id"]) + (b.get("sort") or 0,))
    for r_idx, blk in enumerate(block_list, start=2):
        chain = _chain_of(blk["id"])
        c_id = ws3.cell(row=r_idx, column=1, value=blk["id"])
        c_id.font = ID_FONT
        for i, v in enumerate(chain, start=2):
            ws3.cell(row=r_idx, column=i, value=v)
        ws3.cell(row=r_idx, column=8, value=blk.get("context_passage") or "")
        candidates = blk.get("context_passage_candidates")
        ws3.cell(row=r_idx, column=9,
                 value=json.dumps(candidates, ensure_ascii=False, indent=2) if candidates else "")
        for c in range(1, len(ba_headers) + 1):
            ws3.cell(row=r_idx, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    ba_widths = [26, 16, 16, 16, 20, 12, 12, 60, 60]
    for i, wd in enumerate(ba_widths, start=1):
        ws3.column_dimensions[get_column_letter(i)].width = wd
    ws3.row_dimensions[1].height = 30
    ws3.freeze_panes = "B2"
    ws3.auto_filter.ref = f"A1:{get_column_letter(len(ba_headers))}{len(block_list) + 1}"

    ws2 = wb.create_sheet("Ghi chú")
    ws2.column_dimensions["A"].width = 100
    notes = [
        f"Xuất lúc: {datetime.datetime.now():%Y-%m-%d %H:%M}",
        "",
        "• Sửa tay các cột term/level/pos/ipa/def_en/meaning_vi/sort thoải mái - lưu file lại rồi",
        "  đổi MODE = \"import\" trong sync_vocab.py, chạy lại là đẩy hết lên Supabase.",
        "• Cột 'id': GIỮ NGUYÊN với từ đã có (đừng sửa/xoá) - đây là chìa khoá để tool biết CẬP",
        "  NHẬT đúng từ nào. Để TRỐNG ở 1 dòng MỚI TỰ THÊM -> tool hiểu là TẠO TỪ MỚI.",
        "• Cột 'block_id': BẮT BUỘC điền nếu 'id' để trống (tạo mới) - copy y hệt block_id của",
        "  1 dòng khác đã có sẵn trong CÙNG Block muốn thêm từ vào, đừng tự gõ tay.",
        "• Cột Hub/Notebook/Section/Page/Batch/Block: CHỈ để xem cho dễ đối chiếu, tool KHÔNG",
        "  đọc lại mấy cột này lúc import (sửa cũng không ảnh hưởng gì) - muốn đổi Block của 1",
        "  từ thì sửa đúng cột 'block_id'.",
        "• XOÁ HẲN 1 dòng khỏi Excel KHÔNG xoá từ đó trên Supabase - tool chỉ THÊM/CẬP NHẬT,",
        "  không tự xoá gì cả (an toàn). Muốn xoá từ thật thì vào Supabase Table Editor xoá tay.",
        "",
        "• Sheet 'Bài đọc': 1 dòng/1 BLOCK (khác sheet 'Từ vựng' là 1 dòng/1 TỪ). Sửa cột",
        "  'context_passage' (bài đọc chính) thoải mái. Cột 'context_passage_candidates' là",
        "  DẠNG JSON (danh sách bài đọc ứng cử) - sửa phải giữ ĐÚNG cú pháp JSON (ngoặc vuông",
        "  [...], mỗi bài trong ngoặc kép \"...\", cách nhau dấu phẩy) - sai cú pháp sẽ bị BỎ QUA",
        "  dòng đó lúc import (có in cảnh báo), không làm hỏng dữ liệu cũ trên Supabase.",
        "  Để trống ô này -> KHÔNG đụng gì tới context_passage_candidates cũ (giữ nguyên).",
    ]
    for i, line in enumerate(notes, start=1):
        ws2.cell(row=i, column=1, value=line).font = Font(name="Arial", size=11)

    wb.save(EXCEL_FILE)
    print(f"✅ Đã xuất: {EXCEL_FILE}")


# ==============================================================================
# IMPORT: Excel -> Supabase
# ==============================================================================

def do_import():
    if not EXCEL_FILE.exists():
        sys.exit(f"❌ Không thấy file {EXCEL_FILE} - chạy MODE=\"export\" trước đã.")

    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["Từ vựng"]

    header_row = [c.value for c in ws[1]]

    def _col(name_prefix):
        for i, h in enumerate(header_row):
            if h and str(h).startswith(name_prefix):
                return i
        return None

    idx = {
        "id": _col("id"), "block_id": _col("block_id"),
        "sort": header_row.index("sort"), "term": header_row.index("term"),
        "level": header_row.index("level"), "pos": header_row.index("pos"),
        "ipa": header_row.index("ipa"), "def_en": header_row.index("def_en"),
        "meaning_vi": header_row.index("meaning_vi"),
    }

    rows_out = []
    n_new = n_update = n_skipped = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        term = r[idx["term"]]
        if not term:
            continue  # dòng trống hoàn toàn - bỏ qua

        word_id = r[idx["id"]]
        block_id = r[idx["block_id"]]
        if not word_id:
            if not block_id:
                n_skipped += 1
                print(f"  ⚠️ BỎ QUA '{term}': không có id (từ mới) NHƯNG cũng thiếu block_id "
                      "- không biết thêm vào Block nào.")
                continue
            word_id = str(uuid.uuid4())
            n_new += 1
        else:
            n_update += 1

        rows_out.append({
            "id": word_id,
            "block_id": block_id,
            "sort": r[idx["sort"]],
            "term": term,
            "level": r[idx["level"]],
            "pos": r[idx["pos"]],
            "ipa": r[idx["ipa"]],
            "def_en": r[idx["def_en"]],
            "meaning_vi": r[idx["meaning_vi"]],
        })

    print(f"Đọc được {len(rows_out)} dòng hợp lệ ({n_update} cập nhật, {n_new} tạo mới, "
          f"{n_skipped} bị bỏ qua vì thiếu dữ liệu).")
    if not rows_out:
        print("Không có gì để đẩy lên - dừng.")
        return

    headers = dict(HEADERS)
    headers["Prefer"] = "resolution=merge-duplicates,return=minimal"
    for i in range(0, len(rows_out), IMPORT_BATCH_SIZE):
        batch = rows_out[i:i + IMPORT_BATCH_SIZE]
        r = requests.post(f"{REST_URL}/words", headers=headers, json=batch, timeout=60)
        if not r.ok:
            sys.exit(f"❌ Lỗi đẩy lô {i}-{i + len(batch)}: {r.status_code} {r.text}")
        print(f"  -> Đã đẩy {i + len(batch)}/{len(rows_out)} dòng...")

    print("✅ Xong (từ vựng)! Đã đồng bộ lên Supabase.")

    do_import_passages(wb)


def do_import_passages(wb):
    """Đẩy sheet 'Bài đọc' (context_passage/context_passage_candidates của
    TỪNG BLOCK) lên Supabase - sheet này không bắt buộc phải có (file Excel
    xuất từ bản cũ trước khi thêm sheet này vẫn chạy được, chỉ bỏ qua bước
    này). CHỈ PATCH đúng cột có dữ liệu trong Excel - ô để TRỐNG (cả 2 cột)
    thì KHÔNG đụng gì tới giá trị cũ trên Supabase (tránh xoá oan bài đọc
    đã có sẵn chỉ vì Thao không sửa gì ở dòng đó)."""
    if "Bài đọc" not in wb.sheetnames:
        print("(Không thấy sheet 'Bài đọc' trong file - bỏ qua, có thể file xuất từ bản cũ hơn.)")
        return

    ws = wb["Bài đọc"]
    header_row = [c.value for c in ws[1]]

    def _col(name_prefix):
        for i, h in enumerate(header_row):
            if h and str(h).startswith(name_prefix):
                return i
        return None

    idx_id = _col("id")
    idx_passage = _col("context_passage (")
    idx_candidates = _col("context_passage_candidates (")

    n_updated = n_bad_json = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        block_id = row[idx_id]
        if not block_id:
            continue

        body = {}
        passage = row[idx_passage]
        if passage:
            body["context_passage"] = passage
        candidates_raw = row[idx_candidates]
        if candidates_raw:
            try:
                body["context_passage_candidates"] = json.loads(candidates_raw)
            except Exception as e:
                n_bad_json += 1
                print(f"  ⚠️ BỎ QUA context_passage_candidates của block {block_id}: JSON lỗi ({e})")

        if not body:
            continue  # cả 2 cột trống - không có gì để cập nhật, giữ nguyên dữ liệu cũ

        r = requests.patch(f"{REST_URL}/blocks", headers=HEADERS,
                            params={"id": f"eq.{block_id}"}, json=body, timeout=30)
        if not r.ok:
            print(f"  ⚠️ Lỗi cập nhật bài đọc block {block_id}: {r.status_code} {r.text}")
            continue
        n_updated += 1

    extra = f", {n_bad_json} bị bỏ qua vì JSON lỗi" if n_bad_json else ""
    print(f"✅ Xong (bài đọc)! Đã cập nhật {n_updated} block{extra}.")


# ==============================================================================
# CHẠY
# ==============================================================================

def run():
    print(f"📡 Supabase: {SUPABASE_URL}  —  MODE = {MODE}")
    if MODE == "export":
        do_export()
        if AUTO_OPEN_EXCEL:
            try:
                import os
                os.startfile(str(EXCEL_FILE))
            except Exception as e:
                print(f"  (Không tự mở được Excel: {e})")
    elif MODE == "import":
        do_import()
    else:
        sys.exit('❌ MODE phải là "export" hoặc "import".')


if __name__ == "__main__":
    run()
